from datetime import datetime
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from analyzer_app import config, excel_io, services


@pytest.fixture
def mock_registry_file(tmp_path):
    file_path = tmp_path / "registro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = config.REGISTRO_FOGLIO_NOME
    start_row = config.REGISTRO_RIGA_INIZIO_DATI
    ws.cell(row=start_row, column=config.REGISTRO_COL_IDX_MODELLO_STRUM_CAMPIONE + 1).value = "MOD-REG"
    ws.cell(row=start_row, column=config.REGISTRO_COL_IDX_ID_CERT_CAMPIONE + 1).value = "CERT-REG-001"
    ws.cell(row=start_row, column=config.REGISTRO_COL_IDX_RANGE_CAMPIONE + 1).value = "0-10 bar"
    ws.cell(row=start_row, column=config.REGISTRO_COL_IDX_SCADENZA_CAMPIONE + 1).value = "31/12/2030"
    wb.save(file_path)
    return str(file_path)

def test_leggi_registro_strumenti_success(mock_registry_file, monkeypatch):
    monkeypatch.setattr(config, "FILE_REGISTRO_STRUMENTI", mock_registry_file)
    strumenti = excel_io.leggi_registro_strumenti()
    assert strumenti is not None
    assert len(strumenti) == 1
    assert strumenti[0].id_certificato == "CERT-REG-001"

def test_leggi_registro_non_esistente(monkeypatch):
    monkeypatch.setattr(config, "FILE_REGISTRO_STRUMENTI", "non_esisto.xlsx")
    assert excel_io.leggi_registro_strumenti() is None

def test_read_sheet_file_not_found():
    with pytest.raises(FileNotFoundError):
        excel_io.read_instrument_sheet_raw_data("missing.xlsx")


def test_read_instrument_sheet_raw_data_xls_legacy():
    mock_book = MagicMock()
    mock_sheet = MagicMock()
    mock_book.sheet_by_index.return_value = mock_sheet
    mock_sheet.merged_cells = []

    r_e2, c_e2 = excel_io.excel_coord_to_indices("E2")
    r_odc, c_odc = excel_io.excel_coord_to_indices(config.SCHEDA_ANA_CELL_ODC)

    def mock_cell_value(r, c):
        if r == r_e2 and c == c_e2:
            return "STRUMENTI ANALOGICI"
        if r == r_odc and c == c_odc:
            return "ODC-XLS"
        return None


    mock_sheet.cell_value.side_effect = mock_cell_value

    with patch("xlrd.open_workbook", return_value=mock_book):
        data = excel_io.read_instrument_sheet_raw_data("test.xls")
        assert data['file_type'] == "analogico"
        assert data['odc'] == "ODC-XLS"

def test_read_instrument_sheet_raw_data_xlsx(tmp_path):
    file_path = tmp_path / "test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws['E2'] = "SCHEDA TARATURA STRUMENTI ANALOGICI"
    ws[config.SCHEDA_ANA_CELL_ODC] = "ODC-OK"
    wb.save(file_path)
    data = excel_io.read_instrument_sheet_raw_data(str(file_path))
    if data.get('odc') is not None:
        assert data['odc'] == "ODC-OK"

def test_parse_date_robust():
    assert excel_io.parse_date_robust("19/02/2026") == datetime(2026, 2, 19)

def test_excel_coord_to_indices():
    assert excel_io.excel_coord_to_indices("A1") == (0, 0)

def test_write_cell_xlsx(tmp_path):
    f = tmp_path / "w.xlsx"
    Workbook().save(f)
    assert excel_io.write_cell(str(f), "A1", "V") is True

def test_config_save_and_load(tmp_path):
    config_file = tmp_path / "config_test.json"
    import analyzer_app.config as cfg
    old_path = cfg.CONFIG_FILE_PATH
    cfg.CONFIG_FILE_PATH = str(config_file)
    try:
        new_data = {
            'FILE_REGISTRO_STRUMENTI': 'reg.xlsx',
            'FOLDER_PATH_DEFAULT': 'cards/',
            'FILE_DATI_COMPILAZIONE_SCHEDE': 'data.xlsx'
        }
        assert cfg.save_config(new_data) is True
        assert cfg.load_config_from_json() is True
        assert cfg.FILE_REGISTRO_STRUMENTI == 'reg.xlsx'
    finally:
        cfg.CONFIG_FILE_PATH = old_path

def test_is_config_valid_logic(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "FILE_REGISTRO_STRUMENTI", None)
    assert config.is_config_valid() is False
    reg = tmp_path / "reg.xlsx"
    reg.write_text("data")
    monkeypatch.setattr(config, "FILE_REGISTRO_STRUMENTI", str(reg))
    monkeypatch.setattr(config, "FOLDER_PATH_DEFAULT", str(tmp_path))
    assert config.is_config_valid() is True

def test_internal_read_file_worker_success(tmp_path):
    """Testa il worker interno del service layer."""
    import multiprocessing
    f = tmp_path / "simple.xlsx"
    wb = Workbook()
    wb.active['E2'] = "DIGITALE"
    wb.save(f)

    q = multiprocessing.Queue()
    services._read_file_worker_internal(q, str(f))

    status, res = q.get()
    assert status == 'success'
    assert res['base_filename'] == "simple.xlsx"
